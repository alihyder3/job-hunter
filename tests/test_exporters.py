"""Tests for data exporters."""

import os
import csv
from pathlib import Path
from openpyxl import load_workbook
from job_hunter.models import JobPosting
from job_hunter.exporters import DataExporter


def test_save_as_csv(tmp_path: Path) -> None:
    """Test CSV export functionality."""
    output_file = tmp_path / "test_output.csv"
    
    postings = [
        JobPosting(
            employer_name="CompanyA",
            role_title="Developer",
            role_details="Build apps",
            application_url="https://example.com/a"
        ),
        JobPosting(
            employer_name="CompanyB",
            role_title="Designer",
            role_details="Create designs",
            application_url="https://example.com/b",
            submission_deadline="2026-06-01"
        )
    ]
    
    DataExporter.save_as_csv(postings, str(output_file))
    
    assert output_file.exists()
    
    with open(output_file, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        rows = list(reader)
    
    assert len(rows) == 2
    assert rows[0]["Company"] == "CompanyA"
    assert rows[0]["Title"] == "Developer"
    assert rows[0]["Deadline"] == "Not specified"
    assert rows[1]["Company"] == "CompanyB"
    assert rows[1]["Deadline"] == "2026-06-01"


def test_save_as_xlsx(tmp_path: Path) -> None:
    """Test Excel export functionality."""
    output_file = tmp_path / "test_output.xlsx"
    
    postings = [
        JobPosting(
            employer_name="FirmX",
            role_title="Analyst",
            role_details="Analyze trends",
            application_url="https://example.com/x"
        ),
        JobPosting(
            employer_name="FirmY",
            role_title="Manager",
            role_details="Manage projects",
            application_url="https://example.com/y",
            submission_deadline="2026-09-15"
        )
    ]
    
    DataExporter.save_as_xlsx(postings, str(output_file))
    
    assert output_file.exists()
    
    workbook = load_workbook(output_file)
    worksheet = workbook.active
    
    assert worksheet.title == "Job Opportunities"
    assert worksheet.cell(1, 1).value == "Company"
    assert worksheet.cell(1, 2).value == "Title"
    assert worksheet.cell(1, 3).value == "Description"
    assert worksheet.cell(1, 4).value == "Apply Link"
    assert worksheet.cell(1, 5).value == "Deadline"
    
    assert worksheet.cell(2, 1).value == "FirmX"
    assert worksheet.cell(2, 2).value == "Analyst"
    assert worksheet.cell(3, 1).value == "FirmY"
    assert worksheet.cell(3, 5).value == "2026-09-15"


def test_csv_creates_parent_directories(tmp_path: Path) -> None:
    """Test that CSV export creates parent directories."""
    output_file = tmp_path / "subdir" / "nested" / "output.csv"
    
    postings = [
        JobPosting(
            employer_name="TestCo",
            role_title="Tester",
            role_details="Test software",
            application_url="https://example.com"
        )
    ]
    
    DataExporter.save_as_csv(postings, str(output_file))
    
    assert output_file.exists()
    assert output_file.parent.exists()


def test_xlsx_creates_parent_directories(tmp_path: Path) -> None:
    """Test that Excel export creates parent directories."""
    output_file = tmp_path / "another" / "path" / "output.xlsx"
    
    postings = [
        JobPosting(
            employer_name="OrgZ",
            role_title="Writer",
            role_details="Write content",
            application_url="https://example.com"
        )
    ]
    
    DataExporter.save_as_xlsx(postings, str(output_file))
    
    assert output_file.exists()
    assert output_file.parent.exists()
