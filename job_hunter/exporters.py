"""Data export utilities for job postings."""

import csv
from pathlib import Path
from typing import List
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]

from .models import JobPosting


class DataExporter:
    """Exports job postings to various formats."""
    
    COLUMN_HEADERS = ["Company", "Title", "Description", "Apply Link", "Deadline"]
    
    @staticmethod
    def save_as_csv(postings: List[JobPosting], output_path: str) -> None:
        """Export job postings to CSV file."""
        filepath = Path(output_path)
        filepath.parent.mkdir(parents=True, exist_ok=True)
        
        with open(filepath, 'w', newline='', encoding='utf-8') as csv_file:
            csv_writer = csv.DictWriter(csv_file, fieldnames=DataExporter.COLUMN_HEADERS)
            csv_writer.writeheader()
            
            for posting in postings:
                csv_writer.writerow(posting.to_row_dict())
        
        print(f"CSV exported: {filepath.absolute()}")
    
    @staticmethod
    def save_as_xlsx(postings: List[JobPosting], output_path: str) -> None:
        """Export job postings to Excel file."""
        filepath = Path(output_path)
        filepath.parent.mkdir(parents=True, exist_ok=True)
        
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Job Opportunities"
        
        header_style = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")
        
        for col_idx, header_name in enumerate(DataExporter.COLUMN_HEADERS, start=1):
            cell = worksheet.cell(row=1, column=col_idx, value=header_name)
            cell.font = header_style
            cell.fill = header_fill
        
        for row_idx, posting in enumerate(postings, start=2):
            row_data = posting.to_row_dict()
            for col_idx, header_name in enumerate(DataExporter.COLUMN_HEADERS, start=1):
                worksheet.cell(row=row_idx, column=col_idx, value=row_data[header_name])
        
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
        
        workbook.save(filepath)
        print(f"Excel exported: {filepath.absolute()}")
